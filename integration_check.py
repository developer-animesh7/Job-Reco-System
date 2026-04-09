import io
import json
import os

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "resume_ai_project.settings")

import django
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from docx import Document


def build_docx_bytes():
    buffer = io.BytesIO()
    doc = Document()
    doc.add_heading("Animesh Kumar", level=1)
    doc.add_paragraph("Python Developer")
    doc.add_paragraph("Skills: Python, Django, SQL, Docker, AWS")
    doc.add_paragraph("Experience: Built REST APIs and backend services")
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def main():
    django.setup()

    username = "integration_user"
    password = "pass12345"

    user_model = get_user_model()
    user, created = user_model.objects.get_or_create(username=username)
    if created:
        user.set_password(password)
        user.save()

    client = Client()
    logged_in = client.login(username=username, password=password)
    if not logged_in:
        print("LOGIN_FAILED")
        return 1

    file_content = build_docx_bytes()
    upload = SimpleUploadedFile(
        "resume.docx",
        file_content,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    response = client.post("/analyze-resume/", {"file": upload})

    print("STATUS", response.status_code)
    try:
        payload = response.json()
    except Exception:
        print("JSON_PARSE_FAILED")
        print(response.content[:300])
        return 1

    print("KEYS", sorted(list(payload.keys())))
    if "skills" in payload:
        print("SKILLS_COUNT", len(payload.get("skills") or []))
    if "recommendations" in payload:
        print("RECOMMENDATIONS_COUNT", len(payload.get("recommendations") or []))

    if response.status_code not in (200, 201):
        print("PAYLOAD", json.dumps(payload)[:400])
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
